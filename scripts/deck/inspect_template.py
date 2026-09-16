"""Dump the structure of the company PowerPoint template.

Read-only diagnostic. Reports slide size, layout placeholders, the theme's font
and colour schemes, and where each media file is referenced, so the deck
generator can be written against what the template actually contains.

Usage:
    python scripts/deck/inspect_template.py docs/presentation/company-template.pptx
"""

from __future__ import annotations

import io
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.util import Emu

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"a": A, "r": R}

# The Windows console defaults to a legacy code page that cannot encode CJK
# typeface names, which are common in Chinese templates.
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")


def cm(value: int | None) -> str:
    return "-" if value is None else f"{Emu(value).cm:.2f}"


def report_slides(prs: Presentation) -> None:
    print("=" * 72)
    print("PRESENTATION")
    print("=" * 72)
    print(f"slide size : {cm(prs.slide_width)}cm x {cm(prs.slide_height)}cm")
    ratio = prs.slide_width / prs.slide_height
    label = "16:9" if abs(ratio - 16 / 9) < 0.02 else f"ratio {ratio:.4f}"
    print(f"aspect     : {label}")
    print(f"slides     : {len(prs.slides)}")
    print(f"layouts    : {len(prs.slide_layouts)}")


def report_layouts(prs: Presentation) -> None:
    print()
    print("=" * 72)
    print("SLIDE LAYOUTS")
    print("=" * 72)
    for i, layout in enumerate(prs.slide_layouts):
        print(f"[{i}] {layout.name!r}")
        for shape in layout.placeholders:
            ph = shape.placeholder_format
            idx = "-" if ph.idx is None else ph.idx
            print(
                f"    idx={idx} type={ph.type} name={shape.name!r} "
                f"pos=({cm(shape.left)},{cm(shape.top)}) "
                f"size=({cm(shape.width)}x{cm(shape.height)})"
            )
        if not layout.placeholders:
            print("    (no placeholders)")
        extra = [s.name for s in layout.shapes if not s.is_placeholder]
        if extra:
            print(f"    non-placeholder shapes: {extra}")
        print()


def report_existing_slides(prs: Presentation) -> None:
    print("=" * 72)
    print("EXISTING SLIDES")
    print("=" * 72)
    for i, slide in enumerate(prs.slides):
        print(f"[{i}] layout={slide.slide_layout.name!r}")
        for shape in slide.shapes:
            kind = "placeholder" if shape.is_placeholder else shape.shape_type
            text = ""
            if shape.has_text_frame:
                raw = shape.text_frame.text.replace("\n", " | ")
                text = f" text={raw[:90]!r}"
            print(f"    {shape.name!r} ({kind}){text}")
        print()


def report_fonts(prs: Presentation) -> None:
    print("=" * 72)
    print("EXPLICITLY SET FONTS (slides + layouts)")
    print("=" * 72)
    fonts: set[str] = set()
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        if run.font.name:
                            fonts.add(run.font.name)
    for layout in prs.slide_layouts:
        for shape in layout.placeholders:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        if run.font.name:
                            fonts.add(run.font.name)
    print("  " + "\n  ".join(sorted(fonts)) if fonts else "  (none; theme fonts apply)")


def report_theme(path: Path) -> None:
    print()
    print("=" * 72)
    print("THEME")
    print("=" * 72)
    with zipfile.ZipFile(path) as z:
        for name in sorted(n for n in z.namelist() if n.startswith("ppt/theme/")):
            root = ET.fromstring(z.read(name))
            print(f"--- {name} ---")
            scheme = root.find(".//a:fontScheme", NS)
            if scheme is not None:
                print(f"  fontScheme: {scheme.get('name')}")
                for role in ("majorFont", "minorFont"):
                    el = scheme.find(f"a:{role}", NS)
                    latin = el.find("a:latin", NS)
                    ea = el.find("a:ea", NS)
                    cs = el.find("a:cs", NS)
                    print(
                        f"    {role}: latin={latin.get('typeface')!r} "
                        f"ea={ea.get('typeface')!r} cs={cs.get('typeface')!r}"
                    )
                    for f in el.findall("a:font", NS):
                        typeface = f.get("typeface")
                        if typeface:
                            print(f"      script={f.get('script')!r} {typeface!r}")
            colors = root.find(".//a:clrScheme", NS)
            if colors is not None:
                print(f"  clrScheme: {colors.get('name')}")
                for child in colors:
                    c = child[0]
                    val = c.get("val") or c.get("lastClr")
                    print(f"    {child.tag.split('}')[-1]:<8} {val}")


def report_media(path: Path) -> None:
    print()
    print("=" * 72)
    print("MEDIA AND WHERE IT IS USED")
    print("=" * 72)
    with zipfile.ZipFile(path) as z:
        media = [n for n in z.namelist() if n.startswith("ppt/media/")]
        if not media:
            print("  (no media in package)")
            return
        for m in sorted(media):
            print(f"  {m}  {z.getinfo(m).file_size:,} bytes")

        # Map every media target back to the part that refers to it.
        print()
        print("  references:")
        for rels in sorted(n for n in z.namelist() if n.endswith(".rels")):
            base = rels.rsplit("/_rels/", 1)[0] if "/_rels/" in rels else "."
            root = ET.fromstring(z.read(rels))
            for rel in root:
                target = rel.get("Target", "")
                if "media/" not in target:
                    continue
                owner = f"{base}/{Path(rels).name.replace('.rels', '')}"
                print(f"    {owner}  ->  {target}  ({rel.get('Type','').split('/')[-1]})")


def main(path_str: str) -> int:
    path = Path(path_str)
    if not path.exists():
        print(f"missing: {path}")
        return 1
    prs = Presentation(str(path))
    report_slides(prs)
    report_layouts(prs)
    report_existing_slides(prs)
    report_fonts(prs)
    report_theme(path)
    report_media(path)
    return 0


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else "docs/presentation/company-template.pptx"
    raise SystemExit(main(target))
